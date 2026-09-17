"""
importers.py — Logique d'import en masse des étudiants.
"""

import csv
import io
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from django.db import transaction

from .models import (
    Eleve, DossierAcademique, DossierMilitaire, DossierSante,
    ContactParent, Hebergement,
)


# ─── Constantes ───────────────────────────────────────────────────────────────

VALID_SEXE      = {'H', 'M', 'HOMME', 'MASCULIN', 'F', 'FEMME', 'FEMININ'}
VALID_SERIE_BAC = {'C', 'D', 'TMGM', 'TSGM', 'LM', 'LO', 'ETRANGERE'}
VALID_SEMESTRE  = {'S1', 'S2', 'S3', 'S4', 'S5', 'S6', 'S7', 'S8'}
VALID_VOIE      = {'1', '2', '3', '4'}
VALID_NIVEAU    = {'3', '4', '4-DD', '4-E', '5-DD'}

# Codes département valides (depuis DossierAcademique.CHOIX_DEPARTEMENT)
# departement est un CharField stockant le code court, pas une FK.
DEPT_MAP = {
    # codes directs
    'IRT': 'IRT', 'SID': 'SID', 'GE': 'GE',
    'GM': 'GM', 'GC': 'GC-HE', 'GC-HE': 'GC-HE', 'MPG': 'MPG',
    # noms longs (normalisés sans accents)
    'INFORMATIQUE RESEAUX ET TELECOMMUNICATIONS': 'IRT',
    'INFORMATIQUE RESEAU ET TELECOMMUNICATION': 'IRT',
    'STATISTIQUE ET INGENIERIE DES DONNEES': 'SID',
    'STATISTIQUE INGENIERIE DES DONNEES': 'SID',
    'GENIE ELECTRIQUE': 'GE',
    'GENIE MECANIQUE': 'GM',
    'GENIE CIVIL': 'GC-HE',
    'MINES PETROLE ET GAZ': 'MPG',
    'MINE PETROLE ET GAZ': 'MPG',
    # alias courts avec tirets/espaces
    'GC HE': 'GC-HE',
}

NIVEAU_ALIASES = {
    '3E ANNEE': '3', '3EME': '3', '3': '3',
    '4E ANNEE': '4', '4EME': '4', '4': '4',
    '4E ANNEE DOUBLE DIPLOME': '4-DD', '4DD': '4-DD', '4-DD': '4-DD', '4E DD': '4-DD',
    '5E E': '4-E', '5E': '4-E', '4E ANNEE ECHANGE': '4-E', '4E': '4-E', '4-E': '4-E',
    '5E ANNEE DOUBLE DIPLOME': '5-DD', '5DD': '5-DD', '5-DD': '5-DD', '5E DD': '5-DD',
}


# ─── Utilitaires ──────────────────────────────────────────────────────────────

def _norm(v):
    import unicodedata
    s = unicodedata.normalize('NFD', str(v or '')).encode('ascii', 'ignore').decode()
    return s.strip().upper()


def _parse_date(raw):
    if not raw:
        return None
    if isinstance(raw, (date, datetime)):
        return raw.strftime('%Y-%m-%d')
    s = str(raw).strip()
    for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y', '%Y/%m/%d', '%d.%m.%Y'):
        try:
            return datetime.strptime(s, fmt).strftime('%Y-%m-%d')
        except ValueError:
            continue
    return None


def _parse_decimal(raw, default=None):
    try:
        return Decimal(str(raw).replace(',', '.').strip())
    except (InvalidOperation, TypeError, ValueError):
        return default


def _norm_sexe(v):
    n = _norm(v)
    return 'F' if n.startswith('F') else 'H'


def _norm_niveau(v):
    n = _norm(v).replace(' ', '').replace('-', '')
    for alias, code in NIVEAU_ALIASES.items():
        if _norm(alias).replace(' ', '').replace('-', '') == n:
            return code
    return ''


def _norm_voie(v):
    s = str(v or '').strip()
    m = re.search(r'[1-4]', s)
    return m.group() if m else '1'


def _resolve_dept(raw):
    """Résout une valeur brute vers le code département valide (str) ou None."""
    if not raw:
        return None
    n = _norm(raw)
    # Essai direct
    if n in DEPT_MAP:
        return DEPT_MAP[n]
    # Essai en enlevant les tirets/espaces
    n2 = n.replace('-', '').replace(' ', '')
    for key, code in DEPT_MAP.items():
        if key.replace('-', '').replace(' ', '') == n2:
            return code
    return None


# ─── Lecteurs de fichier ──────────────────────────────────────────────────────

def read_excel(file_obj):
    """
    Read dossier Excel. Returns (rows, meta).

    meta = {sheet_name, format} where format is always 'dossier'.
    Raises ValueError for liste définitive 3A/4A, legacy templates, or bad headers.
    """
    try:
        import openpyxl
    except ImportError:
        raise ImportError("openpyxl est requis. pip install openpyxl")

    from etudiants.importers_dossier import (
        detect_dossier_format,
        find_dossier_header_row,
        looks_like_liste_or_legacy,
    )

    wb = openpyxl.load_workbook(file_obj, data_only=True)
    # Prefer sheet named Etudiants when present
    sheet_name = 'Etudiants' if 'Etudiants' in wb.sheetnames else wb.active.title
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError(
            'Fichier Excel vide. Téléchargez le modèle dossier Polyspace.'
        )

    found = find_dossier_header_row(rows)
    if not found:
        # Probe first non-empty row for clear reject messages
        probe = []
        for row in rows[:10]:
            cells = [str(c).strip() if c is not None else '' for c in row]
            if any(cells):
                probe = cells
                break
        hint = looks_like_liste_or_legacy(probe)
        if hint == 'liste_definitive':
            raise ValueError(
                'Format liste définitive 3A/4A non pris en charge. '
                'Téléchargez le modèle dossier Polyspace (colonnes matricule, nom_famille, …).'
            )
        if hint == 'legacy':
            raise ValueError(
                'Ancien modèle d’import non pris en charge. '
                'Téléchargez le modèle dossier Polyspace via Import → Télécharger le modèle.'
            )
        raise ValueError(
            'En-têtes dossier introuvables. '
            'Téléchargez le modèle Polyspace (ligne d’en-têtes snake_case).'
        )

    header_idx, headers = found
    fmt = detect_dossier_format(headers)
    result = []
    for row in rows[header_idx + 1:]:
        if all(v is None or str(v).strip() == '' for v in row):
            continue
        row_dict = dict(zip(headers, row))
        # Skip empty data area + hidden dropdown ref lists written below the table
        mat = row_dict.get('matricule')
        if mat is None or str(mat).strip() == '':
            continue
        mat_s = str(mat).strip()
        if mat_s.startswith('_'):
            break  # ref lists start with _pays / _wilayas / …
        digits = re.sub(r'\D', '', mat_s)
        if not digits:
            continue
        result.append(row_dict)
    return result, {'sheet_name': sheet_name, 'format': fmt}


def read_csv(file_obj):
    raw = file_obj.read()
    for enc in ('utf-8-sig', 'utf-8', 'latin-1'):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError("Encodage CSV non reconnu. Sauvegardez en UTF-8.")
    sample = text[:2048]
    sep = ';' if sample.count(';') >= sample.count(',') else ','
    reader = csv.DictReader(io.StringIO(text), delimiter=sep)
    return [dict(row) for row in reader]


# ─── Mapping ligne → champs Django ───────────────────────────────────────────

class RowMapper:
    def _pick(self, row, *keys):
        normed = {_norm(k): v for k, v in row.items()}
        for key in keys:
            v = normed.get(_norm(key))
            if v is not None and str(v).strip() != '':
                return str(v).strip()
        return ''

    def map(self, row, row_num):
        errors = []

        def err(field, msg):
            errors.append({'row': row_num, 'field': field, 'message': msg})

        p = self._pick

        # ── Eleve ─────────────────────────────────────────────────────────────
        matricule_raw = p(row, 'matricule', 'Matricule', 'mat')
        try:
            matricule = int(re.sub(r'\D', '', matricule_raw))
        except (ValueError, TypeError):
            matricule = None
            err('matricule', f'Valeur invalide : "{matricule_raw}"')

        nom = p(row, 'nom_famille', 'nom', 'Nom', 'last_name')
        if not nom:
            err('nom_famille', 'Champ obligatoire manquant')

        prenom = p(row, 'prenom', 'Prénom', 'prenoms', 'first_name')
        if not prenom:
            err('prenom', 'Champ obligatoire manquant')

        nni = p(row, 'nni', 'NNI', 'identifiant_national')
        if not nni:
            nni = f'AUTO-{matricule_raw or row_num}'

        date_naissance = _parse_date(p(row, 'date_naissance', 'naissance', 'ddn'))
        if not date_naissance:
            date_naissance = '2000-01-01'

        date_inscription = _parse_date(p(row, 'date_premiere_inscription', 'date_inscription'))
        if not date_inscription:
            date_inscription = f'{date.today().year}-09-01'

        moyenne_bac = _parse_decimal(p(row, 'moyenne_bac', 'moyenne', 'moy_bac'), Decimal('10.00'))
        if moyenne_bac < 0 or moyenne_bac > 20:
            err('moyenne_bac', f'Doit être entre 0 et 20 (reçu : {moyenne_bac})')
            moyenne_bac = Decimal('10.00')

        serie_raw = _norm(p(row, 'serie_bac', 'serie', 'filiere_bac'))
        serie_bac = serie_raw if serie_raw in VALID_SERIE_BAC else 'C'

        categorie_bac_raw = _norm(p(row, 'categorie_bac', 'cat_bac'))
        categorie_bac = 'Etranger' if 'ETRANGER' in categorie_bac_raw else 'National'

        voie_acces = _norm_voie(p(row, 'voie_acces', 'voie', 'acces'))

        eleve_data = dict(
            matricule=matricule,
            num_bac=p(row, 'num_bac', 'numero_bac', 'bac') or '—',
            nni=nni,
            sexe=_norm_sexe(p(row, 'sexe', 'genre', 'sex')),
            prenom=prenom,
            nom_famille=nom,
            date_naissance=date_naissance,
            lieu_naissance=p(row, 'lieu_naissance', 'lieu', 'ville_naissance') or '—',
            nationalite=p(row, 'nationalite', 'nationalité', 'nationality') or 'Mauritanienne',
            categorie_bac=categorie_bac,
            serie_bac=serie_bac,
            moyenne_bac=moyenne_bac,
            ecole_bac=p(row, 'ecole_bac', 'lycee', 'etablissement_bac') or '—',
            date_premiere_inscription=date_inscription,
            voie_acces=voie_acces,
            diplome_acces=p(row, 'diplome_acces', 'diplome') or '—',
            etablissement_diplome=p(row, 'etablissement_diplome', 'etablissement') or '',
            adresse_primaire=p(row, 'adresse_primaire', 'adresse') or '—',
            adresse_secondaire=p(row, 'adresse_secondaire') or '',
            resident_avec_parents=True,
            compte_bankily=p(row, 'compte_bankily', 'bankily') or '',
            email_perso=p(row, 'email_perso', 'email', 'mail') or f'{matricule_raw or row_num}@import.esp.mr',
            tel1=p(row, 'telephone', 'tel', 'tel1', 'phone') or '00000000',
            tel2_whatsapp=p(row, 'whatsapp', 'tel2', 'tel_whatsapp') or '',
        )

        # ── DossierAcademique ─────────────────────────────────────────────────
        dept_raw = p(row, 'departement', 'filiere', 'département', 'dept')
        dept_code = _resolve_dept(dept_raw)
        if dept_code is None:
            if dept_raw:
                err('departement', f'Département inconnu : "{dept_raw}". Valeurs : IRT, GE, GM, GC-HE, SID, MPG.')
            dept_code = 'IRT'  # valeur par défaut non-fatale

        niveau_raw = p(row, 'niveau', 'annee', 'année', 'niveau_actuel')
        niveau = _norm_niveau(niveau_raw)
        if not niveau:
            err('niveau', f'Niveau inconnu : "{niveau_raw}". Valeurs : 3, 4, 4-DD, 4-E, 5-DD.')
            niveau = '3'

        semestre_raw = _norm(p(row, 'semestre', 'semestre_actuel'))
        semestre = semestre_raw if semestre_raw in VALID_SEMESTRE else 'S1'

        da_data = dict(
            departement=dept_code,       # CharField, pas une FK
            niveau_actuel=niveau,
            semestre_actuel=semestre,
            donnees_semestres={},
            parcours=p(row, 'parcours', 'statut') or 'En cours normal',
            diplome=p(row, 'diplome_obtenu') or '',
            etablissement_echange=p(row, 'etablissement_echange') or '',
            etablissement_double_diplome=p(row, 'etablissement_double_diplome') or '',
            specialite_mobilite=p(row, 'specialite_mobilite') or '',
        )

        # ── DossierSante (optionnel) ──────────────────────────────────────────
        poids = _parse_decimal(p(row, 'poids', 'poids_kg'))
        taille = _parse_decimal(p(row, 'taille', 'taille_cm'))
        gs_raw = _norm(p(row, 'groupe_sanguin', 'groupe_sang', 'blood'))
        gs = gs_raw if gs_raw in {'A+', 'A-', 'B+', 'B-', 'AB+', 'AB-', 'O+', 'O-'} else None

        ds_data = dict(
            groupe_sanguin=gs or 'O+',
            assureur=p(row, 'assureur') or '',
            num_assure=p(row, 'num_assure', 'numero_assure') or '',
            antecedents_medicaux=p(row, 'antecedents', 'antecedents_medicaux') or '',
            maladies_chroniques=p(row, 'maladies_chroniques') or '',
            medicaments_a_vie=p(row, 'medicaments') or '',
            poids_kg=poids,
            taille_cm=taille,
        ) if (gs or poids or taille) else None

        # ── ContactParent (optionnel) ─────────────────────────────────────────
        tel_urgence = p(row, 'tel_urgence', 'telephone_urgence', 'tel_pere', 'tel_parent')
        cp_data = dict(
            prenom_pere=p(row, 'prenom_pere', 'nom_pere') or '',
            nom_famille_pere=p(row, 'nom_famille_pere') or nom,
            fonction_pere=p(row, 'fonction_pere') or '',
            tel_pere=p(row, 'tel_pere', 'telephone_pere') or '',
            prenom_mere=p(row, 'prenom_mere', 'nom_mere') or '',
            nom_famille_mere=p(row, 'nom_famille_mere') or '',
            tel_mere=p(row, 'tel_mere', 'telephone_mere') or '',
            nom_urgence=p(row, 'nom_urgence', 'contact_urgence') or '',
            tel_urgence=tel_urgence or '00000000',
        ) if tel_urgence else None

        return {
            'eleve': eleve_data,
            'dossier_academique': da_data,
            'dossier_militaire': _auto_dossier_militaire(niveau, dept_code),
            'dossier_sante': ds_data,
            'contact_parent': cp_data,
        }, errors


def _auto_dossier_militaire(niveau, dept_code):
    from etudiants.importers_liste_definitive import dossier_militaire_from_niveau_dept
    return dossier_militaire_from_niveau_dept(niveau or '3', dept_code or '')


# ─── Import principal ─────────────────────────────────────────────────────────

class BulkImporter:
    def __init__(self):
        self.mapper = RowMapper()

    def run(self, rows, meta=None):
        from etudiants.importers_dossier import FORMAT_DOSSIER, DossierRowMapper

        meta = meta or {}
        fmt = meta.get('format') or FORMAT_DOSSIER
        if fmt != FORMAT_DOSSIER:
            raise ValueError(
                'Format d’import non pris en charge. '
                'Utilisez uniquement le modèle dossier Polyspace.'
            )

        dossier_mapper = DossierRowMapper()
        created = 0
        skipped = 0
        all_errors = []

        for idx, raw_row in enumerate(rows):
            row_num = idx + 2
            mapped, field_errors = dossier_mapper.map(raw_row, row_num)
            fatal_fields = DossierRowMapper.FATAL

            if field_errors:
                fatal = [e for e in field_errors if e['field'] in fatal_fields]
                all_errors.extend(field_errors)
                if fatal:
                    skipped += 1
                    continue

            try:
                with transaction.atomic():
                    self._create_row(mapped)
                created += 1
            except Exception as exc:
                skipped += 1
                msg = str(exc)
                if 'matricule' in msg and 'unique' in msg.lower():
                    msg = f"Matricule {mapped['eleve'].get('matricule')} déjà existant."
                elif 'nni' in msg and 'unique' in msg.lower():
                    msg = f"NNI {mapped['eleve'].get('nni')} déjà existant."
                all_errors.append({'row': row_num, 'field': 'db', 'message': msg})

        return {'created': created, 'skipped': skipped, 'errors': all_errors[:200]}

    def _create_row(self, mapped):
        eleve = Eleve.objects.create(**dict(mapped['eleve']))

        da = dict(mapped['dossier_academique'])
        da['eleve'] = eleve
        DossierAcademique.objects.create(**da)

        dm = mapped.get('dossier_militaire')
        if dm and (dm.get('compagnie') or dm.get('section')):
            dm_payload = {
                'eleve': eleve,
                'compagnie': dm.get('compagnie') or '',
                'section': dm.get('section') or '',
                'sport_pratique': dm.get('sport_pratique') or '',
            }
            for field in (
                'tour_poitrine',
                'tour_ceinture',
                'tour_taille',
                'tour_bassin',
                'tour_cou',
                'longueur_manche',
                'longueur_dos',
                'longueur_cote',
                'pointure',
            ):
                if dm.get(field) is not None:
                    dm_payload[field] = dm[field]
            DossierMilitaire.objects.create(**dm_payload)

        if mapped.get('dossier_sante'):
            ds = dict(mapped['dossier_sante'])
            if not ds.get('poids_kg') or not ds.get('taille_cm'):
                ds['poids_kg'] = None
                ds['taille_cm'] = None
            ds['eleve'] = eleve
            DossierSante.objects.create(**ds)

        if mapped.get('contact_parent'):
            cp = dict(mapped['contact_parent'])
            cp['eleve'] = eleve
            ContactParent.objects.create(**cp)

        if mapped.get('hebergement'):
            heb = dict(mapped['hebergement'])
            heb['eleve'] = eleve
            Hebergement.objects.create(**heb)


# ─── Générateur de templates ──────────────────────────────────────────────────

def generate_excel_template():
    """Serve the unified dossier workbook (same as Desktop preview)."""
    from etudiants.dossier_excel_template import workbook_to_bytes
    return workbook_to_bytes()


def generate_csv_template():
    """Header-only CSV from dossier COLUMN_HEADERS."""
    from etudiants.dossier_excel_template import COLUMN_HEADERS
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=';')
    writer.writerow(COLUMN_HEADERS)
    return buf.getvalue()
