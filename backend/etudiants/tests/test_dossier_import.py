"""Dossier Excel import — sole accepted format."""

from io import BytesIO
from pathlib import Path

from django.test import TestCase
from openpyxl import Workbook, load_workbook

from etudiants.dossier_excel_template import COLUMN_HEADERS, workbook_to_bytes
from etudiants.importers import BulkImporter, generate_csv_template, generate_excel_template, read_excel
from etudiants.importers_dossier import (
    FORMAT_DOSSIER,
    DossierRowMapper,
    detect_dossier_format,
    find_dossier_header_row,
)
from etudiants.models import ContactParent, DossierAcademique, DossierMilitaire, Eleve, Hebergement

FIXTURES = Path(__file__).resolve().parent / 'fixtures'


def _dossier_workbook_bytes(extra_rows=None):
    """Minimal in-memory dossier xlsx (banner + headers + data)."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Etudiants'
    ws.append(['Polyspace — modèle dossier élève'])
    ws.append(list(COLUMN_HEADERS))
    row = {h: '' for h in COLUMN_HEADERS}
    row.update({
        'matricule': 251280,
        'nom_famille': 'Ould Ahmed',
        'prenom': 'Mohamed',
        'nni': '9800123456',
        'sexe': 'M',
        'date_naissance': '15/05/2002',
        'nationalite': 'Mauritanie',
        'wilaya_naissance': 'Nouakchott',
        'commune_naissance': 'Tevragh-Zeina',
        'num_bac': '12345',
        'categorie_bac': 'National',
        'serie_bac': 'C',
        'moyenne_bac': '14,50',
        'ecole_bac': 'Lycée Nationale',
        'departement': 'IRT',
        'niveau': '3e année',
        'statut_academique': 'Normal',
        'annee_univ_1ere': '2024-2025',
        'date_premiere_inscription': '01/09/2024',
        'voie_acces': '1',
        'diplome_acces': 'CNIM',
        'etablissement_diplome': 'IPGEI',
        'adresse_primaire': 'Tevragh-Zeina',
        'telephone': '31234567',
        'email_perso': 'med.ahmed@gmail.com',
        'resident_avec_parents': 'Oui',
        'prenom_pere': 'Ahmed',
        'nom_famille_pere': 'Ould Ahmed',
        'tel_urgence': '20001122',
        'groupe_sanguin': 'O+',
        'poids_kg': '72',
        'taille_cm': '178',
        'sport_pratique': 'Football',
        'batiment': 'R1',
        'etage': '2',
        'aile': 'A',
        'chambre': '205',
        'lit': 'B',
    })
    ws.append([row[h] for h in COLUMN_HEADERS])
    if extra_rows:
        for er in extra_rows:
            r = {h: '' for h in COLUMN_HEADERS}
            r.update(er)
            ws.append([r[h] for h in COLUMN_HEADERS])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


class DossierDetectTests(TestCase):
    def test_detect_ok(self):
        self.assertEqual(detect_dossier_format(list(COLUMN_HEADERS)), FORMAT_DOSSIER)

    def test_detect_rejects_liste(self):
        h = [
            'N°', 'Matricule', 'Nom et Prénom', 'Département', 'Voie', 'S',
            'Institut', 'NNI',
        ]
        with self.assertRaises(ValueError) as ctx:
            detect_dossier_format(h)
        self.assertIn('liste définitive', str(ctx.exception).lower())

    def test_find_header_skips_banner(self):
        rows = [
            ['Polyspace banner'],
            list(COLUMN_HEADERS),
            [251280, 'X', 'Y'],
        ]
        idx, headers = find_dossier_header_row(rows)
        self.assertEqual(idx, 1)
        self.assertEqual(headers[0], 'matricule')


class DossierReadRejectTests(TestCase):
    def test_reject_3a_fixture(self):
        path = FIXTURES / 'liste_definitive_3a_sample.xlsx'
        with path.open('rb') as fh:
            with self.assertRaises(ValueError) as ctx:
                read_excel(fh)
        self.assertIn('liste définitive', str(ctx.exception).lower())

    def test_reject_4a_fixture(self):
        path = FIXTURES / 'liste_definitive_4a_sample.xlsx'
        with path.open('rb') as fh:
            with self.assertRaises(ValueError) as ctx:
                read_excel(fh)
        self.assertIn('liste définitive', str(ctx.exception).lower())


class DossierImportTests(TestCase):
    def test_import_creates_full_row(self):
        data = _dossier_workbook_bytes()
        rows, meta = read_excel(BytesIO(data))
        self.assertEqual(meta['format'], FORMAT_DOSSIER)
        self.assertEqual(len(rows), 1)

        report = BulkImporter().run(rows, meta=meta)
        self.assertEqual(report['created'], 1, report)
        self.assertEqual(report['skipped'], 0, report)

        e = Eleve.objects.get(matricule=251280)
        self.assertEqual(e.prenom, 'Mohamed')
        self.assertEqual(e.nom_famille, 'Ould Ahmed')
        self.assertEqual(e.voie_acces, '1')
        self.assertFalse(e.profil_incomplet)
        self.assertEqual(e.nni, '9800123456')
        self.assertIn('Tevragh', e.lieu_naissance)

        da = DossierAcademique.objects.get(eleve=e)
        self.assertEqual(da.departement, 'IRT')
        self.assertEqual(da.niveau_actuel, '3')
        self.assertEqual(da.parcours, 'En cours normal')

        dm = DossierMilitaire.objects.get(eleve=e)
        self.assertEqual(dm.compagnie, '1re Compagnie')
        self.assertEqual(dm.section, 'Section 11')
        self.assertEqual(dm.sport_pratique, 'Football')

        self.assertTrue(ContactParent.objects.filter(eleve=e).exists())
        self.assertTrue(Hebergement.objects.filter(eleve=e).exists())

    def test_redoublant_not_voie(self):
        extra = [{
            'matricule': 251281,
            'nom_famille': 'Ba',
            'prenom': 'Aicha',
            'sexe': 'F',
            'departement': 'GE',
            'niveau': '4e année',
            'statut_academique': 'Redoublant',
            'voie_acces': '',
            'date_naissance': '01/01/2001',
        }]
        data = _dossier_workbook_bytes(extra_rows=extra)
        rows, meta = read_excel(BytesIO(data))
        # only use second data row
        report = BulkImporter().run([rows[1]], meta=meta)
        self.assertEqual(report['created'], 1, report)
        e = Eleve.objects.get(matricule=251281)
        self.assertEqual(e.voie_acces, '')
        self.assertEqual(e.dossier_academique.parcours, 'Redoublant')
        self.assertEqual(e.dossier_academique.niveau_actuel, '4')

    def test_3e_requires_voie(self):
        mapped, errors = DossierRowMapper().map({
            'matricule': '100',
            'nom_famille': 'X',
            'prenom': 'Y',
            'departement': 'IRT',
            'niveau': '3e année',
            'voie_acces': '',
            'sexe': 'M',
        }, 2)
        self.assertTrue(any(e['field'] == 'voie_acces' for e in errors))


class DossierTemplateTests(TestCase):
    def test_excel_template_has_dossier_headers(self):
        content = generate_excel_template()
        wb = load_workbook(BytesIO(content), data_only=True)
        self.assertIn('Etudiants', wb.sheetnames)
        ws = wb['Etudiants']
        rows = list(ws.iter_rows(values_only=True))
        found = find_dossier_header_row(rows)
        self.assertIsNotNone(found)
        _, headers = found
        for core in ('matricule', 'nom_famille', 'prenom', 'departement', 'niveau', 'statut_academique'):
            self.assertIn(core, headers)

    def test_workbook_to_bytes_is_dossier(self):
        content = generate_excel_template()
        self.assertTrue(content.startswith(b'PK'))
        self.assertGreater(len(content), 1000)
        # Same builder path as Desktop preview
        self.assertTrue(callable(workbook_to_bytes))

    def test_csv_template_headers(self):
        csv_text = generate_csv_template()
        first = csv_text.strip().splitlines()[0]
        for h in ('matricule', 'nom_famille', 'statut_academique'):
            self.assertIn(h, first)
